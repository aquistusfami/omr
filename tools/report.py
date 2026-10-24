#!/usr/bin/env python3
"""Build the marking report from the hub's data folder.

    python3 tools/report.py                  # data/omr -> bao-cao.xlsx
    python3 tools/report.py --dir data/omr -o /tmp/bao-cao.xlsx

Scores are recomputed from key.json (written by the app when the answer key is
set), so fixing a key and re-running this is enough to rescore everything.
openpyxl is the only dependency: uv run --with openpyxl python3 tools/report.py
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
P1_COUNT = 10
P2_FIRST = 11
P2_COUNT = 5
LETTERS = "abcdef"


def load(dir_path: Path) -> tuple[list[dict], dict]:
    rows = []
    sheet_file = dir_path / "sheets.jsonl"
    if sheet_file.exists():
        for line in sheet_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue  # a torn line writes nothing; the phone resends it
    key_file = dir_path / "key.json"
    key = json.loads(key_file.read_text(encoding="utf-8")) if key_file.exists() else {}
    return rows, key


def score(row: dict, key: dict) -> tuple[float, int, int] | None:
    """Same rule as web/src/key.ts: P1 full credit; P2 all six or weighted."""
    cfg = key.get("cfg") or {}
    keys = key.get("keys") or {}
    answers = keys.get(row.get("examCode") or "")
    if not answers or not row.get("ok"):
        return None

    w1 = float(cfg.get("p1", 2))
    w2 = float(cfg.get("p2", 2))
    split = cfg.get("mode") == "split"

    correct_p1 = 0
    for i in range(P1_COUNT):
        want = answers["p1"][i] if i < len(answers.get("p1", [])) else ""
        got = (row.get("p1") or {}).get(str(i + 1), "")
        if want and got == want:
            correct_p1 += 1

    part_two = 0.0
    correct_p2 = 0
    for q in range(P2_COUNT):
        want = answers["p2"][q] if q < len(answers.get("p2", [])) else ""
        got = (row.get("p2") or {}).get(str(P2_FIRST + q), "")
        matched = sum(1 for i in range(6) if i < len(want) and i < len(got) and want[i] == got[i])
        if len(want) == 6 and matched == 6:
            part_two += w2
            correct_p2 += 1
        elif split:
            part_two += w2 * matched / 6

    return round(correct_p1 * w1 + part_two, 2), correct_p1, correct_p2


def table(rows: list[dict], key: dict) -> tuple[list[str], list[list]]:
    head = ["stt", "thoi_gian", "ma_sv", "ma_de"]
    head += [f"cau{i}" for i in range(1, P1_COUNT + 1)]
    head += [f"{P2_FIRST + q}{LETTERS[i]}" for q in range(P2_COUNT) for i in range(6)]
    head += ["diem", "dung_p1", "dung_p2", "trung", "can_kiem_tra"]

    body = []
    for n, row in enumerate(rows, 1):
        s = score(row, key)
        line = [
            n,
            datetime.fromtimestamp((row.get("at") or 0) / 1000).strftime("%Y-%m-%d %H:%M:%S"),
            row.get("studentId", ""),
            row.get("examCode", ""),
        ]
        line += [(row.get("p1") or {}).get(str(i), "") for i in range(1, P1_COUNT + 1)]
        for q in range(P2_COUNT):
            got = (row.get("p2") or {}).get(str(P2_FIRST + q), "")
            line += [c if c not in ("", "?") else "" for c in (got + "------")[:6]]
        line += [
            s[0] if s else "",
            s[1] if s else "",
            s[2] if s else "",
            "x" if row.get("duplicate") else "",
            "" if row.get("ok") else f"REJECT:{row.get('reason', '')}",
        ]
        if row.get("ok") and row.get("flags"):
            line[-1] = " ".join(row["flags"])
        body.append(line)
    return head, body


def write_xlsx(path: Path, rows: list[dict], key: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    head, body = table(rows, key)
    scored = [s for s in (score(r, key) for r in rows) if s]
    maximum = 10 * float((key.get("cfg") or {}).get("p1", 2)) + 5 * float((key.get("cfg") or {}).get("p2", 2))

    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sach"
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for line in body:
        ws.append(line)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14 if col > 4 else 12

    hist = wb.create_sheet("Phong diem")
    hist.append(["diem", "so phieu", "ti le"])
    for cell in hist[1]:
        cell.font = Font(bold=True)
    if scored:
        marks = [s[0] for s in scored]
        top = int(maximum) + 1
        for mark in range(top):
            n = sum(1 for m in marks if mark <= m < mark + 1)
            if n:
                hist.append([f"{mark}-{mark + 1}", n, round(n / len(marks), 3)])
        hist.append([])
        hist.append(["so phieu co diem", len(marks)])
        hist.append(["diem trung binh", round(statistics.mean(marks), 2)])
        hist.append(["trung vi", statistics.median(marks)])
        hist.append(["cao nhat", max(marks)])
        hist.append(["thap nhat", min(marks)])
        hist.append(["thang diem", maximum])

    code = wb.create_sheet("Theo ma de")
    code.append(["ma de", "so phieu", "diem TB", "cao nhat", "thap nhat", "so o can xem"])
    for cell in code[1]:
        cell.font = Font(bold=True)
    for exam in sorted({r.get("examCode", "") for r in rows}):
        group = [r for r in rows if r.get("examCode") == exam]
        marks = [s[0] for s in (score(r, key) for r in group) if s]
        code.append(
            [
                exam or "?",
                len(group),
                round(statistics.mean(marks), 2) if marks else "",
                max(marks) if marks else "",
                min(marks) if marks else "",
                sum(1 for r in group if r.get("flags") or not r.get("ok")),
            ]
        )

    sus = wb.create_sheet("Nghi van")
    sus.append(["ma_sv", "ma_de", "ly do", "anh"])
    for cell in sus[1]:
        cell.font = Font(bold=True)
    for r in rows:
        if r.get("ok") and not r.get("flags") and not r.get("duplicate"):
            continue
        reasons = []
        if not r.get("ok"):
            reasons.append(f"tu choi: {r.get('reason', '')}")
        reasons += r.get("flags") or []
        if r.get("duplicate"):
            reasons.append("trung ma SV")
        sus.append([r.get("studentId", ""), r.get("examCode", ""), ", ".join(reasons), f"thumbs/{r.get('uid')}.jpg"])

    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Tao bao cao Excel tu thu muc du lieu cua hub")
    ap.add_argument("--dir", default=str(ROOT / "data" / "omr"), help="thu muc du lieu (mac dinh data/omr)")
    ap.add_argument("-o", "--out", default=None, help="file xlsx dau ra")
    args = ap.parse_args()

    data = Path(args.dir).expanduser().resolve()
    if not data.exists():
        print(f"Khong tim thay thu muc: {data}", file=sys.stderr)
        return 2
    rows, key = load(data)
    if not rows:
        print(f"Chua co phieu nao trong {data}/sheets.jsonl", file=sys.stderr)
        return 2

    out = Path(args.out) if args.out else data / "bao-cao.xlsx"
    try:
        write_xlsx(out, rows, key)
    except ImportError:
        print("Thieu openpyxl. Chay lai bang:\n  uv run --with openpyxl python3 tools/report.py", file=sys.stderr)
        return 2
    scored = sum(1 for r in rows if score(r, key))
    print(f"{len(rows)} phieu, {scored} co diem -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
